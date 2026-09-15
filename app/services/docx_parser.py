"""
Extracts raw text from DOCX resumes using python-docx.

Ported from the original hybrid resume-parser project's
`extractors/file_parser.py` (DOCX half).
"""
from __future__ import annotations

import io
import logging
from pathlib import Path

import docx

from app.services.pdf_parser import FileParsingError, _normalize_whitespace

logger = logging.getLogger(__name__)


class UnsupportedFileTypeError(Exception):
    pass


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file, including paragraphs and table cells."""
    try:
        document = docx.Document(io.BytesIO(file_bytes))
        parts: list[str] = []

        for para in document.paragraphs:
            if para.text.strip():
                parts.append(para.text)

        # Resumes often use tables for layout (e.g. skills grids, contact blocks)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))

        text = "\n".join(parts)
        if not text.strip():
            raise FileParsingError("No extractable text found in DOCX file.")
        return _normalize_whitespace(text)
    except FileParsingError:
        raise
    except Exception as exc:  # noqa: BLE001
        logger.exception("Failed to parse DOCX")
        raise FileParsingError(f"Could not parse DOCX file: {exc}") from exc


def extract_text(filename: str, file_bytes: bytes) -> str:
    """Dispatch to the correct extractor (PDF/DOCX) based on file extension."""
    from app.services.pdf_parser import extract_text_from_pdf

    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_bytes)
    if ext == ".docx":
        return extract_text_from_docx(file_bytes)
    raise UnsupportedFileTypeError(f"Unsupported file type '{ext}'. Only .pdf and .docx are supported.")
